"""
pptx_reader.py - PowerPoint (.pptx) -> Claude content blocks.

For every slide it extracts: title + all text (including inside groups and tables),
speaker notes, and embedded pictures (diagrams / screenshots / scanned slide images),
so image-only slides are still readable by Claude's vision.

Install:  pip install python-pptx pillow
Old .ppt files are not supported - re-save them as .pptx in PowerPoint.
"""
import base64
import hashlib
import io
import logging

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

try:
    from PIL import Image
except ImportError:  # Pillow is optional - only needed for odd formats / big images
    Image = None

log = logging.getLogger(__name__)

MAX_SLIDES = 80
MAX_IMAGES = 25                 # total pictures sent per deck (cost control)
MIN_IMAGE_BYTES = 8_000         # skip tiny icons/bullets
MAX_IMAGE_BYTES = 4_500_000     # Claude limit is 5 MB per image
NATIVE_TYPES = {"image/jpeg", "image/png", "image/gif", "image/webp"}


def _picture_block(blob: bytes, content_type: str, seen: set) -> dict | None:
    """Return an image content block, or None if the picture should be skipped."""
    if len(blob) < MIN_IMAGE_BYTES:
        return None
    digest = hashlib.md5(blob).hexdigest()
    if digest in seen:  # logos / backgrounds repeated on every slide
        return None
    seen.add(digest)

    media_type = (content_type or "").lower().replace("image/jpg", "image/jpeg")
    if media_type not in NATIVE_TYPES or len(blob) > MAX_IMAGE_BYTES:
        if Image is None:
            return None
        try:  # convert bmp/tiff/etc. to PNG, shrink very large images to JPEG
            img = Image.open(io.BytesIO(blob))
            img.load()
            if max(img.size) > 2000:
                img.thumbnail((2000, 2000))
            buf = io.BytesIO()
            if len(blob) > MAX_IMAGE_BYTES:
                img.convert("RGB").save(buf, "JPEG", quality=85)
                media_type = "image/jpeg"
            else:
                img.save(buf, "PNG")
                media_type = "image/png"
            blob = buf.getvalue()
            if len(blob) > MAX_IMAGE_BYTES:
                return None
        except Exception:  # noqa: BLE001  (emf/wmf and other unsupported formats)
            return None

    return {
        "type": "image",
        "source": {"type": "base64", "media_type": media_type, "data": base64.standard_b64encode(blob).decode()},
    }


def _walk(shapes):
    """Yield every shape, descending into groups."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)
        else:
            yield shape


def pptx_to_content_blocks(pptx_path: str) -> tuple[list[dict], dict]:
    """
    Returns (content_blocks, stats).
    stats: {"slides": n, "images": n, "truncated": bool}
    """
    prs = Presentation(pptx_path)
    blocks: list[dict] = []
    seen: set = set()
    images_sent = 0
    total = len(prs.slides)

    for i, slide in enumerate(prs.slides, start=1):
        if i > MAX_SLIDES:
            break

        title = ""
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text_frame.text.strip()
        except Exception:  # noqa: BLE001
            pass

        text_parts: list[str] = []
        pictures: list[dict] = []
        for shape in _walk(slide.shapes):
            try:
                if shape.has_text_frame:
                    t = shape.text_frame.text.strip()
                    if t and t != title:
                        text_parts.append(t)
                elif getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                        text_parts.append(" | ".join(cells))
                elif getattr(shape, "has_chart", False) and shape.has_chart:
                    text_parts.append("[chart on this slide]")
                if shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.PLACEHOLDER):
                    try:
                        img = shape.image
                    except Exception:  # noqa: BLE001  (placeholder without a picture)
                        continue
                    if images_sent + len(pictures) < MAX_IMAGES:
                        blk = _picture_block(img.blob, img.content_type, seen)
                        if blk:
                            pictures.append(blk)
            except Exception:  # noqa: BLE001 - one odd shape must not kill the whole deck
                log.warning("Skipping unreadable shape on slide %d", i, exc_info=True)

        notes = ""
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:  # noqa: BLE001
            pass

        header = f"--- Slide {i}" + (f": {title}" if title else "") + " ---"
        body = "\n".join(text_parts)
        if notes:
            body += f"\n[Speaker notes] {notes}"
        if not body.strip() and not pictures:
            body = "(empty slide)"
        blocks.append({"type": "text", "text": f"{header}\n{body}".strip()})
        blocks.extend(pictures)
        images_sent += len(pictures)

    stats = {"slides": total, "images": images_sent, "truncated": total > MAX_SLIDES}
    log.info("PPTX %s: %s", pptx_path, stats)
    return blocks, stats
